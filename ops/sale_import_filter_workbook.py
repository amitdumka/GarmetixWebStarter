import argparse
import re
from pathlib import Path

import openpyxl


def norm(value):
    return re.sub(r"\s+", " ", str(value or "").strip()).upper()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source")
    parser.add_argument("--sources", nargs="+")
    parser.add_argument("--output", required=True)
    parser.add_argument("--invoices", nargs="+", required=True)
    args = parser.parse_args()

    source_paths = [Path(item) for item in (args.sources or ([args.source] if args.source else []))]
    if not source_paths:
        raise SystemExit("Pass --source or --sources.")

    invoices = {item.strip() for item in args.invoices if item.strip()}
    wb = openpyxl.load_workbook(source_paths[0])
    for sheet_name, header_row in [("Sale Report", 4), ("Item Details", 3)]:
        ws = wb[sheet_name]
        headers = {norm(cell.value): idx for idx, cell in enumerate(ws[header_row], start=1) if cell.value is not None}
        invoice_col = headers.get(norm("Invoice No./Txn No."))
        if not invoice_col:
            continue
        for row_idx in range(ws.max_row, header_row, -1):
            invoice = str(ws.cell(row=row_idx, column=invoice_col).value or "").strip()
            if invoice not in invoices:
                ws.delete_rows(row_idx)

    for extra_source in source_paths[1:]:
        extra = openpyxl.load_workbook(extra_source, data_only=False)
        for sheet_name, header_row in [("Sale Report", 4), ("Item Details", 3)]:
            if sheet_name not in extra.sheetnames or sheet_name not in wb.sheetnames:
                continue
            src_ws = extra[sheet_name]
            dst_ws = wb[sheet_name]
            headers = {norm(cell.value): idx for idx, cell in enumerate(src_ws[header_row], start=1) if cell.value is not None}
            invoice_col = headers.get(norm("Invoice No./Txn No."))
            if not invoice_col:
                continue
            for row in src_ws.iter_rows(min_row=header_row + 1):
                invoice = str(row[invoice_col - 1].value or "").strip()
                if invoice not in invoices:
                    continue
                values = [cell.value for cell in row]
                dst_ws.append(values)

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    print(output)


if __name__ == "__main__":
    main()
