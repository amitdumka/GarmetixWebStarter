import argparse
import csv
import json
import mimetypes
import os
import re
import sys
import uuid
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from urllib.parse import urlencode
from urllib.request import Request, urlopen
from urllib.error import HTTPError

import openpyxl


COMPANY_ID = "2d337c03-fe27-4f7c-a70f-9b3a4f482c24"
STORE_GROUP_ID = "6be6141f-a4e2-4579-abdb-232de5b00784"
STORE_ID = "6bb60584-9c2b-4af3-ad9f-8ecce968b723"

SOURCE_WORKBOOK = Path(r"C:\AIArea\SaleReport_01_08_25_to_31_03_26_FIN YEAR_2025_2026.xlsx")
INVOICES = ["AF/2025/1345", "AF/2025/1365", "AF/2025/1380", "AF/2025/1382", "AF/2025/1386"]
SERVICE_INVOICE = "AF/2025/1365"

LINE_ACTIONS = {
    ("AF/2025/1345", "", "SKAM"): {
        "action": "create_stock_product",
        "overrideBarcode": "SKAM0001",
        "createProductAndStock": True,
        "note": "Amit approved create stock/product for SKAM with barcode SKAM0001.",
    },
    ("AF/2025/1345", "", "Juti(Gift)"): {
        "action": "map_existing_barcode",
        "overrideBarcode": "25105059",
        "createProductAndStock": False,
        "note": "Amit supplied existing barcode 25105059.",
    },
    ("AF/2025/1380", "", "Arvind Shirting TALP52216-1"): {
        "action": "create_stock_product",
        "overrideBarcode": "TALP52216-1",
        "createProductAndStock": True,
        "note": "Amit approved create stock/product with opening bridge qty same as billed.",
    },
    ("AF/2025/1380", "M20905101008", "Arvind TAS Suiting Black"): {
        "action": "map_existing_barcode",
        "overrideBarcode": "M2090501008",
        "createProductAndStock": False,
        "note": "Amit corrected source barcode to existing M2090501008.",
    },
    ("AF/2025/1382", "r24067142", "Arvind Mills Shirting r24067142"): {
        "action": "create_stock_product",
        "overrideBarcode": "R24067142",
        "createProductAndStock": True,
        "note": "Amit approved create stock/product with opening bridge qty same as billed.",
    },
    ("AF/2025/1386", "10120100012", "Arvind Mill Shirting White R240337046"): {
        "action": "map_existing_barcode",
        "overrideBarcode": "R240337046",
        "createProductAndStock": False,
        "note": "Amit approved mapping to existing R240337046.",
    },
}

SERVICE_LINES = [
    {
        "serviceCode": "38639007445",
        "name": "Tailoring 3Pcs Suit",
        "quantity": 1,
        "customerRate": 4800,
        "description": "3 Pcs Suit Stitching",
    },
    {
        "serviceCode": "38611802580",
        "name": "Tailoring Shirt",
        "quantity": 1,
        "customerRate": 450,
        "description": "Shirting Stitching",
    },
]


def norm(value):
    return re.sub(r"\s+", " ", str(value or "").strip()).upper()


class Api:
    def __init__(self, base_url, token=None):
        self.base_url = base_url.rstrip("/")
        self.token = token

    def request(self, method, path, body=None, query=None, headers=None):
        url = f"{self.base_url}/{path.lstrip('/')}"
        if query:
            url += "?" + urlencode({k: v for k, v in query.items() if v is not None})
        data = None
        all_headers = {"Accept": "application/json"}
        if self.token:
            all_headers["Authorization"] = f"Bearer {self.token}"
        if headers:
            all_headers.update(headers)
        if body is not None:
            data = json.dumps(body).encode("utf-8")
            all_headers["Content-Type"] = "application/json"
        req = Request(url, data=data, method=method, headers=all_headers)
        try:
            with urlopen(req, timeout=120) as response:
                text = response.read().decode("utf-8")
                return json.loads(text) if text else None
        except HTTPError as exc:
            detail = exc.read().decode("utf-8", errors="replace")
            raise RuntimeError(f"{method} {url} failed with {exc.code}: {detail}") from exc

    def upload(self, path, file_path, query=None):
        boundary = "----garmetixcodex" + uuid.uuid4().hex
        url = f"{self.base_url}/{path.lstrip('/')}"
        if query:
            url += "?" + urlencode({k: v for k, v in query.items() if v is not None})
        content_type = mimetypes.guess_type(str(file_path))[0] or "application/octet-stream"
        data = file_path.read_bytes()
        parts = [
            f"--{boundary}\r\n".encode(),
            f'Content-Disposition: form-data; name="file"; filename="{file_path.name}"\r\n'.encode(),
            f"Content-Type: {content_type}\r\n\r\n".encode(),
            data,
            b"\r\n",
            f"--{boundary}--\r\n".encode(),
        ]
        headers = {
            "Accept": "application/json",
            "Content-Type": f"multipart/form-data; boundary={boundary}",
        }
        if self.token:
            headers["Authorization"] = f"Bearer {self.token}"
        req = Request(url, data=b"".join(parts), method="POST", headers=headers)
        try:
            with urlopen(req, timeout=180) as response:
                text = response.read().decode("utf-8")
                return json.loads(text) if text else None
        except HTTPError as exc:
            detail = exc.read().decode("utf-8", errors="replace")
            raise RuntimeError(f"POST {url} failed with {exc.code}: {detail}") from exc


def filter_workbook(output_path: Path):
    wb = openpyxl.load_workbook(SOURCE_WORKBOOK)
    for sheet_name, header_row in [("Sale Report", 4), ("Item Details", 3)]:
        ws = wb[sheet_name]
        headers = {norm(cell.value): idx for idx, cell in enumerate(ws[header_row], start=1) if cell.value is not None}
        invoice_col = headers.get(norm("Invoice No./Txn No."))
        if not invoice_col:
            continue
        for row_idx in range(ws.max_row, header_row, -1):
            invoice = str(ws.cell(row=row_idx, column=invoice_col).value or "").strip()
            if invoice not in INVOICES:
                ws.delete_rows(row_idx)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def get_token(api: Api, username: str, password: str) -> str:
    response = api.request("POST", "auth/login", {"userName": username, "password": password})
    token = response.get("token") if isinstance(response, dict) else None
    if not token:
        raise RuntimeError("Login succeeded but token was not returned.")
    return token


def read_rows(value):
    if isinstance(value, list):
        return value
    if isinstance(value, dict):
        for key in ("items", "rows", "data", "results"):
            if isinstance(value.get(key), list):
                return value[key]
    return []


def first_active_salesman(api: Api):
    options = api.request(
        "GET",
        "billing/options",
        query={"companyId": COMPANY_ID, "storeId": STORE_ID, "take": 100},
    )
    salesmen = read_rows(options.get("salesmen") if isinstance(options, dict) else options)
    if not salesmen:
        raise RuntimeError("No active salesman was found for Smart Menswear.")
    return salesmen[0]["id"], salesmen[0].get("name", "Salesman")


def find_imported(api: Api, invoice_no: str):
    result = api.request(
        "GET",
        "sale-import/vyapar/imported",
        query={"companyId": COMPANY_ID, "storeId": STORE_ID, "from": "2026-02-01", "to": "2026-02-28", "q": invoice_no, "pageSize": 50},
    )
    for row in read_rows(result):
        if row.get("invoiceNumber") == invoice_no or row.get("sourceInvoiceNumber") == invoice_no:
            return row
    return None


def find_customer(api: Api, name: str):
    rows = api.request("GET", "customers")
    for row in read_rows(rows):
        if norm(row.get("name")) == norm(name):
            return row
    return None


def ensure_customer(api: Api, name: str, mobile: str, log):
    customer = find_customer(api, name)
    if customer:
        log.append({"action": "reuse_customer", "customer": name, "customerId": customer["id"]})
        return customer
    payload = {
        "companyId": COMPANY_ID,
        "name": name,
        "mobileNumber": mobile,
        "address": "Dumka",
        "city": "Dumka",
        "state": "Jharkhand",
        "country": "India",
    }
    created = api.request("POST", "customers", payload)
    log.append({"action": "create_customer", "customer": name, "customerId": created["id"], "mobileNumber": mobile})
    return created


def ensure_service_item(api: Api, service, log):
    rows = api.request("GET", "tailoring/service-items", query={"storeId": STORE_ID, "activeOnly": "false"})
    for row in read_rows(rows):
        if norm(row.get("serviceCode")) == norm(service["serviceCode"]):
            log.append({"action": "reuse_service_item", "serviceCode": service["serviceCode"], "serviceItemId": row["id"]})
            return row
    payload = {
        "companyId": COMPANY_ID,
        "storeGroupId": STORE_GROUP_ID,
        "storeId": STORE_ID,
        "serviceCode": service["serviceCode"],
        "name": service["name"],
        "category": 0,
        "defaultCustomerRate": service["customerRate"],
        "defaultVendorRate": 0,
        "taxRate": 5,
        "hsnCode": None,
        "productId": None,
        "active": True,
        "remarks": f"Created for Vyapar sale service import {SERVICE_INVOICE}.",
    }
    created = api.request("POST", "tailoring/service-items", payload)
    log.append({"action": "create_service_item", "serviceCode": service["serviceCode"], "serviceItemId": created["id"]})
    return created


def create_service_invoice(api: Api, salesman_id: str, log):
    existing = find_imported(api, SERVICE_INVOICE)
    if existing:
        log.append({"invoice": SERVICE_INVOICE, "action": "skip_service_invoice_exists", "invoiceId": existing["id"], "invoiceNumber": existing["invoiceNumber"]})
        return existing

    customer = ensure_customer(api, "Asish Dubey", "VYAPAR-AF-2025-1365", log)
    service_items = [ensure_service_item(api, service, log) for service in SERVICE_LINES]
    order_payload = {
        "companyId": COMPANY_ID,
        "storeGroupId": STORE_GROUP_ID,
        "storeId": STORE_ID,
        "orderType": 0,
        "customerId": customer["id"],
        "vendorId": None,
        "sourceInvoiceId": None,
        "sourceInvoiceItemId": None,
        "sourceProductId": None,
        "sourceProductName": None,
        "sourceBarcode": None,
        "expectedDeliveryDate": "2026-02-20",
        "measurementsJson": None,
        "customerInstructions": "Imported historical Vyapar tailoring service bill.",
        "internalRemarks": f"VyaparSourceInvoice={SERVICE_INVOICE}; created by Codex batch-01 approved mapping.",
        "lines": [
            {
                "serviceItemId": service_items[index]["id"],
                "serviceName": service["name"],
                "category": 0,
                "garmentName": service["name"],
                "barcode": service["serviceCode"],
                "quantity": service["quantity"],
                "customerRate": service["customerRate"],
                "vendorRate": 0,
                "discountAmount": 0,
                "costResponsibility": 0,
                "expectedDeliveryDate": "2026-02-20",
                "measurementsJson": None,
                "instructions": service["description"],
                "vendorRemarks": None,
            }
            for index, service in enumerate(SERVICE_LINES)
        ],
    }
    order = api.request("POST", "tailoring/orders", order_payload)
    log.append({"invoice": SERVICE_INVOICE, "action": "create_tailoring_order", "orderId": order["id"], "orderNumber": order.get("orderNumber")})

    converted = api.request(
        "POST",
        f"tailoring/orders/{order['id']}/convert-to-service-invoice",
        {
            "invoiceDate": "2026-02-20",
            "salesmanId": salesman_id,
            "additionalPaidAmount": 0,
            "additionalPaymentMode": 0,
            "bankAccountId": None,
            "referenceNumber": "",
            "remarks": f"Converted from Vyapar service bill {SERVICE_INVOICE}.",
        },
    )
    invoice_id = converted["id"]
    log.append({"invoice": SERVICE_INVOICE, "action": "convert_to_service_invoice", "invoiceId": invoice_id, "generatedInvoiceNumber": converted.get("invoiceNumber")})

    updated = api.request(
        "PUT",
        f"billing/sales/{invoice_id}",
        {
            "invoiceNumber": SERVICE_INVOICE,
            "onDate": "2026-02-20",
            "customerName": "Asish Dubey",
            "customerMobileNumber": customer["mobileNumber"],
            "customerGstin": None,
            "salesmanId": salesman_id,
            "remarks": f"VyaparSaleImport; VyaparSourceInvoice={SERVICE_INVOICE}; SourceDate=2026-02-20; Service invoice via Tailoring module; CodexBatch=batch-01.",
        },
    )
    log.append({"invoice": SERVICE_INVOICE, "action": "rename_service_invoice_to_vyapar_number", "invoiceId": invoice_id, "invoiceNumber": updated.get("invoiceNumber")})
    return updated


def apply_line_actions(preview, log):
    confirm_invoices = []
    applied = []
    for invoice in preview.get("invoices", []):
        source_no = invoice.get("sourceInvoiceNumber")
        if source_no == SERVICE_INVOICE:
            continue
        if source_no not in INVOICES:
            continue
        for line in invoice.get("lines", []):
            key = (source_no, line.get("vyaparItemCode") or "", line.get("itemName") or "")
            action = LINE_ACTIONS.get(key)
            if action:
                line["overrideBarcode"] = action["overrideBarcode"]
                line["createProductAndStock"] = action["createProductAndStock"]
                line["importLine"] = True
                applied.append({
                    "invoice": source_no,
                    "lineNumber": line.get("lineNumber"),
                    "itemCode": line.get("vyaparItemCode") or "",
                    "itemName": line.get("itemName") or "",
                    "action": action["action"],
                    "overrideBarcode": action["overrideBarcode"],
                    "createProductAndStock": action["createProductAndStock"],
                    "quantity": line.get("quantity"),
                    "lineTotal": line.get("lineTotal"),
                    "note": action["note"],
                })
            else:
                line["importLine"] = True
        confirm_invoices.append(invoice)
    log.extend({"action": "line_mapping", **row} for row in applied)
    return confirm_invoices, applied


def write_json(path: Path, value):
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False), encoding="utf-8")


def write_csv(path: Path, rows):
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    fieldnames = sorted({key for row in rows for key in row.keys()})
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--base-url", default="http://192.168.11.94:8088/api")
    parser.add_argument("--username", default="garmetix")
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--execute", action="store_true")
    args = parser.parse_args()

    if not args.execute:
        raise SystemExit("Refusing live mutation without --execute.")
    password = os.environ.get("GARMETIX_API_PASSWORD")
    if not password:
        raise SystemExit("Set GARMETIX_API_PASSWORD before running.")

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    log = []
    workbook_path = output_dir / "batch-01-selected-invoices.xlsx"
    filter_workbook(workbook_path)

    api = Api(args.base_url)
    api.token = get_token(api, args.username, password)
    log.append({"action": "login", "username": args.username, "baseUrl": args.base_url, "at": datetime.now().isoformat(timespec="seconds")})

    salesman_id, salesman_name = first_active_salesman(api)
    log.append({"action": "select_salesman", "salesmanId": salesman_id, "salesmanName": salesman_name})

    preview = api.upload(
        "sale-import/vyapar/preview",
        workbook_path,
        query={"companyId": COMPANY_ID, "storeGroupId": STORE_GROUP_ID, "storeId": STORE_ID},
    )
    write_json(output_dir / "batch-01-preview-before-actions.json", preview)
    log.append({
        "action": "preview",
        "invoiceCount": preview.get("invoiceCount"),
        "lineCount": preview.get("lineCount"),
        "matchedLineCount": preview.get("matchedLineCount"),
        "missingLineCount": preview.get("missingLineCount"),
        "warnings": " | ".join(preview.get("warnings", [])),
    })

    create_service_invoice(api, salesman_id, log)

    confirm_invoices, applied = apply_line_actions(preview, log)
    write_json(output_dir / "batch-01-confirm-invoices.json", confirm_invoices)
    confirm_payload = {
        "companyId": COMPANY_ID,
        "storeGroupId": STORE_GROUP_ID,
        "storeId": STORE_ID,
        "useVyaparInvoiceNumbers": True,
        "createMissingProductsAndStock": True,
        "allowStockBridgeForInsufficientStock": True,
        "defaultBankAccountId": None,
        "defaultSalesmanId": salesman_id,
        "paymentBankMappings": [],
        "finalApprovalConfirmed": True,
        "importBatchId": str(uuid.uuid4()),
        "sourceFileName": workbook_path.name,
        "invoices": confirm_invoices,
    }
    write_json(output_dir / "batch-01-confirm-request-redacted.json", {**confirm_payload, "invoices": f"{len(confirm_invoices)} invoice objects saved separately"})
    confirm_result = api.request("POST", "sale-import/vyapar/confirm", confirm_payload)
    write_json(output_dir / "batch-01-confirm-result.json", confirm_result)
    log.append({"action": "sale_import_confirm", **confirm_result})

    verification = []
    for invoice_no in INVOICES:
        row = find_imported(api, invoice_no)
        verification.append({"invoice": invoice_no, "found": bool(row), **(row or {})})
    write_json(output_dir / "batch-01-import-verification.json", verification)
    log.extend({"action": "verify_imported", **row} for row in verification)

    write_json(output_dir / "batch-01-action-log.json", log)
    write_csv(output_dir / "batch-01-action-log.csv", log)

    md = [
        "# Sale Import Batch 01 Action Log",
        "",
        "Backup taken before mutation: `garmetix-srp-db-20260803-225033-IST-SaleImportMissingBatch01-v6.9.43.dump`.",
        "",
        "## Approved Actions Applied",
        "",
    ]
    for row in applied:
        md.append(f"- `{row['invoice']}` | `{row['itemCode']}` | {row['itemName']} -> `{row['overrideBarcode']}` | {row['action']} | qty {row['quantity']} | amount {row['lineTotal']}")
    md.extend(["", "## Service Invoice", ""])
    for item in [entry for entry in log if entry.get("invoice") == SERVICE_INVOICE]:
        md.append(f"- {item.get('action')}: {item}")
    md.extend(["", "## Confirm Result", "", "```json", json.dumps(confirm_result, indent=2), "```", "", "## Verification", ""])
    for row in verification:
        md.append(f"- `{invoice := row['invoice']}`: {'found' if row.get('found') else 'NOT FOUND'} {row.get('invoiceNumber', '')} {row.get('billAmount', '')}")
    (output_dir / "batch-01-action-log.md").write_text("\n".join(md) + "\n", encoding="utf-8")

    print(json.dumps({"confirmResult": confirm_result, "verification": verification, "logPath": str(output_dir / "batch-01-action-log.md")}, indent=2))


if __name__ == "__main__":
    main()
